"""
Export des slides PPTX en images PNG 1080x1920 via Playwright (headless Chromium).
Convertit d'abord le PPTX en HTML puis screenshot chaque slide.

Usage :
    python content/export.py --slides data/processed/arxiv_2401.12345/slides.pptx
"""

import argparse
import asyncio
import json
import logging
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

logger = logging.getLogger(__name__)

SLIDE_W = 1080
SLIDE_H = 1920


def _rgb_to_hex(rgb: RGBColor | None) -> str:
    """Convertit un RGBColor pptx en couleur CSS hex."""
    if rgb is None:
        return "#ffffff"
    return f"#{rgb[0]:02x}{rgb[1]:02x}{rgb[2]:02x}"


def _fill_css(shape) -> str:
    """Retourne la couleur de fond CSS d'un shape pptx."""
    try:
        if shape.fill.type is not None:
            return f"background:{_rgb_to_hex(shape.fill.fore_color.rgb)};"
    except Exception:
        pass
    return ""


def _border_css(shape) -> str:
    """Retourne le CSS de bordure d'un shape pptx."""
    try:
        w = shape.line.width
        if w and w > 0:
            color = _rgb_to_hex(shape.line.color.rgb)
            px = max(1, int(w / 12700))
            return f"border:{px}px solid {color};"
    except Exception:
        pass
    return ""


def _text_html(shape) -> str:
    """Retourne le HTML du contenu texte d'un shape."""
    if not shape.has_text_frame:
        return ""
    parts = []
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            font = run.font
            size = int(font.size / 12700) if font.size else 24
            bold = "bold" if font.bold else "normal"
            try:
                color = _rgb_to_hex(font.color.rgb)
            except Exception:
                color = "#ffffff"
            parts.append(
                f'<span style="font-size:{size}px;font-weight:{bold};color:{color}">'
                f"{run.text}</span>"
            )
    return "".join(parts) or shape.text_frame.text


def _shape_to_html(shape, slide_width: int, slide_height: int) -> str:
    """Convertit un shape pptx en HTML (texte + formes + connecteurs)."""
    import math
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    # ── Connecteur (ligne droite) ──────────────────────────────────────────
    if shape.shape_type == MSO_SHAPE_TYPE.LINE:
        try:
            x1 = shape.left / slide_width * SLIDE_W
            y1 = shape.top / slide_height * SLIDE_H
            x2 = (shape.left + shape.width) / slide_width * SLIDE_W
            y2 = (shape.top + shape.height) / slide_height * SLIDE_H
            dx, dy = x2 - x1, y2 - y1
            length = math.hypot(dx, dy)
            angle = math.degrees(math.atan2(dy, dx))
            try:
                color = _rgb_to_hex(shape.line.color.rgb)
            except Exception:
                color = "#334466"
            try:
                lw = max(1, int(shape.line.width / 12700)) if shape.line.width else 1
            except Exception:
                lw = 1
            return (
                f'<div style="position:absolute;'
                f'left:{x1:.1f}px;top:{y1:.1f}px;'
                f'width:{length:.1f}px;height:{lw}px;'
                f'background:{color};'
                f'transform-origin:0 50%;'
                f'transform:rotate({angle:.2f}deg);"></div>'
            )
        except Exception:
            return ""

    # ── Forme standard (rectangle, ovale, arrondi, textbox) ───────────────
    left_pct = shape.left / slide_width * 100
    top_pct = shape.top / slide_height * 100
    width_pct = shape.width / slide_width * 100
    height_pct = shape.height / slide_height * 100

    style = (
        f"position:absolute;"
        f"left:{left_pct:.3f}%;top:{top_pct:.3f}%;"
        f"width:{width_pct:.3f}%;height:{height_pct:.3f}%;"
    )

    style += _fill_css(shape)
    style += _border_css(shape)

    # Border-radius selon le type de forme
    try:
        from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
        ast = shape.auto_shape_type
        if ast == MSO_AUTO_SHAPE_TYPE.OVAL:
            style += "border-radius:50%;"
        elif ast == MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE:
            style += "border-radius:10%;"
    except Exception:
        pass

    text = _text_html(shape)
    if text:
        style += (
            "display:flex;align-items:center;justify-content:center;"
            "text-align:center;font-family:'Segoe UI',Arial,sans-serif;"
            "line-height:1.3;overflow:hidden;"
        )

    return f'<div style="{style}">{text}</div>'


def pptx_to_html_slides(pptx_path: Path, output_dir: Path) -> list[Path]:
    """
    Convertit chaque slide PPTX en fichier HTML autonome.
    Rend textes, formes géométriques et connecteurs.

    Args:
        pptx_path:  Chemin vers le fichier PPTX
        output_dir: Dossier de destination des HTML

    Returns:
        Liste des fichiers HTML générés (un par slide)
    """
    prs = Presentation(pptx_path)
    output_dir.mkdir(parents=True, exist_ok=True)
    html_files: list[Path] = []

    for i, slide in enumerate(prs.slides):
        try:
            bg_hex = _rgb_to_hex(slide.background.fill.fore_color.rgb)
        except Exception:
            bg_hex = "#0d0d1a"

        shapes_html = [
            _shape_to_html(shape, prs.slide_width, prs.slide_height)
            for shape in slide.shapes
        ]

        html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  * {{ margin:0; padding:0; box-sizing:border-box; }}
  body {{
    width:{SLIDE_W}px; height:{SLIDE_H}px;
    background:{bg_hex};
    position:relative;
    overflow:hidden;
  }}
</style>
</head>
<body>
{"".join(shapes_html)}
</body>
</html>"""

        html_path = output_dir / f"slide_{i:02d}.html"
        html_path.write_text(html, encoding="utf-8")
        html_files.append(html_path)

    logger.info("%d fichiers HTML générés dans %s", len(html_files), output_dir)
    return html_files


async def _screenshot_slides(html_files: list[Path], output_dir: Path) -> list[Path]:
    """Prend un screenshot PNG de chaque fichier HTML via Playwright."""
    from playwright.async_api import async_playwright

    png_files: list[Path] = []

    async with async_playwright() as p:
        browser = await p.chromium.launch()
        page = await browser.new_page(viewport={"width": SLIDE_W, "height": SLIDE_H})

        for html_path in html_files:
            png_path = output_dir / (html_path.stem + ".png")
            await page.goto(f"file:///{html_path.resolve()}")
            await page.screenshot(path=str(png_path), full_page=False)
            png_files.append(png_path)
            logger.info("PNG exporté : %s", png_path.name)

        await browser.close()

    return png_files


def export_slides_to_png(pptx_path: Path) -> list[Path]:
    """
    Pipeline complet : PPTX → HTML → PNG (1080x1920).

    Args:
        pptx_path: Chemin vers le fichier PPTX

    Returns:
        Liste des fichiers PNG générés
    """
    output_dir = pptx_path.parent / "png"
    html_dir = pptx_path.parent / "html"

    html_files = pptx_to_html_slides(pptx_path, html_dir)
    png_files = asyncio.run(_screenshot_slides(html_files, output_dir))

    logger.info("Export terminé : %d PNGs dans %s", len(png_files), output_dir)
    return png_files


def main() -> None:
    """Point d'entrée CLI."""
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s [%(levelname)s] %(name)s : %(message)s",
        datefmt="%Y-%m-%d %H:%M:%S",
    )
    parser = argparse.ArgumentParser(description="Export slides -> PNG 1080x1920 via Playwright")
    parser.add_argument("--slides", required=True,
                        help="Chemin vers slides.pptx ou script.json")
    args = parser.parse_args()

    path = Path(args.slides)
    if not path.exists():
        logger.error("Fichier introuvable : %s", path)
        return

    if path.suffix == ".json":
        # Pipeline direct JSON → HTML (SVG) → PNG
        from content.slides import generate_html_slides
        html_files = generate_html_slides(path)
        output_dir = path.parent / path.stem / "png"
        output_dir.mkdir(parents=True, exist_ok=True)
        png_files = asyncio.run(_screenshot_slides(html_files, output_dir))
    else:
        # Pipeline PPTX → HTML → PNG
        png_files = export_slides_to_png(path)

    print(f"{len(png_files)} images PNG exportées")
    for p in png_files:
        print(f"  {p}")


if __name__ == "__main__":
    main()
